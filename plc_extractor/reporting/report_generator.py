"""
Automated Engineering Report Generator
Creates professional, publication-grade Excel (.xlsx) workbooks and Markdown (.md) reports
detailing PLC profile, hardware specs, network metrics, memory map, and AI integration recommendations.
"""

import logging
import os
import re
import time
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("ReportGenerator")


class ReportGenerator:
    """
    Generates standardized technical integration reports for industrial controllers.
    """

    def __init__(self, output_dir: str = "data/reports"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def _sanitize_name(self, name: str) -> str:
        clean = re.sub(r'[^a-zA-Z0-9_\-\.]', '_', str(name)).strip("._")
        return clean if clean else "Machine-01"

    def generate_excel_report(self, diagnostic_data: Dict[str, Any]) -> str:
        """Generates a styled, multi-sheet Excel report using openpyxl."""
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        station_name = self._sanitize_name(diagnostic_data.get("station_name", "Machine-01"))
        filename = f"Machine_Integration_Report_{station_name}_{timestamp}.xlsx"
        filepath = os.path.join(self.output_dir, filename)

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Style Palettes
        navy_fill = PatternFill(start_color="1F2E4D", end_color="1F2E4D", fill_type="solid")
        cyan_fill = PatternFill(start_color="06B6D4", end_color="06B6D4", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Segoe UI", size=16, bold=True, color="1F2E4D")
        bold_font = Font(name="Segoe UI", size=10, bold=True)
        regular_font = Font(name="Segoe UI", size=10)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        # ----------------------------------------------------------------------
        # SHEET 1: Executive Summary & Controller Baseline
        # ----------------------------------------------------------------------
        ws1 = wb.create_sheet(title="Executive Summary")
        ws1.views.sheetView[0].showGridLines = True

        ws1["A1"] = f"Industrial Machine Integration & Diagnostics Report: {station_name}"
        ws1["A1"].font = title_font
        ws1["A2"] = f"Generated Automatically: {time.strftime('%Y-%m-%d %H:%M:%S')} | Safety Lock: STRICT_READ_ONLY_ENFORCED"
        ws1["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="64748B")

        summary_rows = [
            ("Target IP Address", diagnostic_data.get("target_ip", "127.0.0.1")),
            ("Detected Industrial Protocol", diagnostic_data.get("detected_protocol", "Modbus TCP")),
            ("Controller Brand / Vendor", diagnostic_data.get("vendor", "Siemens / Modbus Compatible")),
            ("Model / Module Type", diagnostic_data.get("module_type", "Standard Industrial PLC")),
            ("Firmware Revision", diagnostic_data.get("firmware_version", "Detected / Compatible")),
            ("Hardware Serial Number", diagnostic_data.get("serial_number", "N/A")),
            ("Physical Rack / Slot", f"Rack {diagnostic_data.get('rack', 0)} / Slot {diagnostic_data.get('slot', 1)}"),
            ("Communication Safety Status", "100% STRICT READ-ONLY (Zero Plant Impact)"),
            ("Network Response Latency", f"{diagnostic_data.get('avg_latency_ms', 0.0):.2f} ms"),
            ("Recommended Polling Frequency", f"{diagnostic_data.get('recommended_poll_rate_hz', 20)} Hz (every {diagnostic_data.get('recommended_interval_ms', 50)} ms)"),
            ("Total Discovered Tags / Registers", diagnostic_data.get("total_tags_count", 0)),
        ]

        ws1.append([])
        ws1.append(["PARAMETER", "VALUE"])
        header_row = ws1.max_row
        for col in range(1, 3):
            cell = ws1.cell(row=header_row, column=col)
            cell.fill = navy_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for param, val in summary_rows:
            ws1.append([param, str(val)])
            r = ws1.max_row
            ws1.cell(row=r, column=1).font = bold_font
            ws1.cell(row=r, column=2).font = regular_font
            ws1.cell(row=r, column=1).border = thin_border
            ws1.cell(row=r, column=2).border = thin_border

        # ----------------------------------------------------------------------
        # SHEET 2: Discovered Memory Map & Tags
        # ----------------------------------------------------------------------
        ws2 = wb.create_sheet(title="Memory Map & Tags")
        ws2.views.sheetView[0].showGridLines = True

        tags_headers = ["Tag / Key Name", "Protocol", "Area / Table", "Address / Offset", "Data Type", "Polling Tier", "Observed Values / Description"]
        ws2.append(tags_headers)
        for col in range(1, len(tags_headers) + 1):
            cell = ws2.cell(row=1, column=col)
            cell.fill = navy_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        tags_list = diagnostic_data.get("tags", [])
        status = diagnostic_data.get("status", "UNKNOWN")

        if not tags_list:
            if status in ("UNREACHABLE", "COMMUNICATION_FAILED", "FAILED"):
                ws2.append(["N/A", diagnostic_data.get("detected_protocol", "N/A"), "N/A", "N/A", "N/A", "N/A", "لا توجد متغيرات مستخرجة - تعذر الاتصال بالمعالج أو لم يستجب للقراءة المباشرة"])
            else:
                ws2.append(["N/A", "N/A", "N/A", "N/A", "N/A", "N/A", "لم يتم اكتشاف متغيرات نشطة"])
            r = ws2.max_row
            for col in range(1, len(tags_headers) + 1):
                ws2.cell(row=r, column=col).border = thin_border
                ws2.cell(row=r, column=col).font = regular_font
        else:
            for t in tags_list:
                ws2.append([
                    t.get("name", "N/A"),
                    t.get("protocol", diagnostic_data.get("detected_protocol", "Modbus TCP")),
                    t.get("area", "Holding Register"),
                    str(t.get("address", "0")),
                    t.get("data_type", "UINT16"),
                    t.get("tier", "Tier 1 (50ms)"),
                    t.get("desc", ""),
                ])
                r = ws2.max_row
                for col in range(1, len(tags_headers) + 1):
                    ws2.cell(row=r, column=col).border = thin_border
                    ws2.cell(row=r, column=col).font = regular_font

        # ----------------------------------------------------------------------
        # SHEET 3: AI & Engineering Integration Recommendations
        # ----------------------------------------------------------------------
        ws3 = wb.create_sheet(title="AI Recommendations")
        ws3.views.sheetView[0].showGridLines = True

        ws3["A1"] = "Optimal Factory Integration & Telemetry Roadmap"
        ws3["A1"].font = title_font

        recommendations = diagnostic_data.get("recommendations", [
            ("Connection Strategy", "Establish a single persistent TCP socket to port 502/102. Avoid socket teardown on loop iterations."),
            ("Sampling Rate", "Set Tier 1 polling to 50ms (20 Hz). This guarantees 10x oversampling for 500ms cycle transitions without CPU starvation."),
            ("Data Persistence", "Stream directly to DuckDB memory buffer with hourly Parquet exports to eliminate SD card wear on edge PCs."),
            ("Cycle Boundaries", "Detect rising edge on Register 0 (Transition from 0 to 1). Confirm with hydraulic pressure threshold (>30 Bar)."),
            ("Safety Guard", "Keep ReadOnlyPLCClient enforced. All write/force methods are strictly jailed at runtime."),
        ])

        ws3.append([])
        ws3.append(["INTEGRATION DOMAIN", "ENGINEERING RECOMMENDATION"])
        header_row = ws3.max_row
        for col in range(1, 3):
            cell = ws3.cell(row=header_row, column=col)
            cell.fill = navy_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for domain, rec in recommendations:
            ws3.append([domain, rec])
            r = ws3.max_row
            ws3.cell(row=r, column=1).font = bold_font
            ws3.cell(row=r, column=2).font = regular_font
            ws3.cell(row=r, column=1).border = thin_border
            ws3.cell(row=r, column=2).border = thin_border

        # Auto-adjust column widths on all sheets
        for ws in wb.worksheets:
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        wb.save(filepath)
        logger.info(f"Excel Integration Report generated successfully: {filepath}")
        return filepath

    def generate_markdown_report(self, diagnostic_data: Dict[str, Any]) -> str:
        """Generates a clean Markdown/Text report ready to upload or review."""
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        station_name = self._sanitize_name(diagnostic_data.get("station_name", "Machine-01"))
        filename = f"Machine_Integration_Report_{station_name}_{timestamp}.md"
        filepath = os.path.join(self.output_dir, filename)

        tags_list = diagnostic_data.get("tags", [])
        status = diagnostic_data.get("status", "UNKNOWN")

        if tags_list:
            tags_table_rows = "\n".join([
                f"| `{t.get('name', 'N/A')}` | `{t.get('data_type', 'UINT16')}` | {t.get('area') or t.get('address', '0')} | {t.get('tier', 'Tier 1 (50ms)')} | {t.get('desc', '')} |"
                for t in tags_list
            ])
            tags_section = f"""| اسم المتغير (Tag Name) | نوع المتغير (Data Type) | العنوان / Offset | معدل التحديث الموصى به | الغرض الصناعي والقيمة الحالية |
| :--- | :--- | :--- | :--- | :--- |
{tags_table_rows}"""
        elif status in ("UNREACHABLE", "COMMUNICATION_FAILED", "FAILED"):
            tags_section = "*لا توجد متغيرات مستخرجة - تعذر الاتصال بالمعالج أو لم يستجب للقراءة المباشرة.*"
        else:
            tags_section = "*لم يتم اكتشاف متغيرات نشطة على المعالج.*"

        rec_items = "\n".join([
            f"{i+1}. **{domain}:**\n   - {rec}"
            for i, (domain, rec) in enumerate(diagnostic_data.get("recommendations", [
                ("طريقة الاتصال المستقر", "الحفاظ على جلسة اتصال TCP واحدة دائمة وتجنب فتح/إغلاق الاتصال في كل لفة."),
                ("محرك التخزين والأداء", "استخدام محرك DuckDB + Parquet المفصول لتجميع البيانات في الذاكرة وكتابتها كل ساعة لمنع استهلاك وسائط التخزين."),
                ("كشف الدورات بدقة", "الربط بين تغير حالة الماكينة وارتفاع ضغط التشغيل لتفادي أي ارتداد إشارة."),
                ("حماية الماكينة بنسبة 100%", "الاستمرار في تفعيل جدار حماية القراءة فقط (ReadOnlyPLCClient) لضمان عدم حدوث أي تغيير في كود الـ PLC الأصلي."),
            ]))
        ])

        md_content = f"""# 🏭 تقرير الفحص الفني الشامل وتوصيات الربط الآمن للماكينة
> **Station:** {station_name} | **Target IP:** `{diagnostic_data.get('target_ip', '127.0.0.1')}`
> **Generated:** {time.strftime('%Y-%m-%d %H:%M:%S')} | **Safety Profile:** `STRICT_READ_ONLY (100% Safe)`
> **Status:** `{status}`

---

## 1. بطاقة تعريف الماكينة والمعالج (Controller Profile)

| البند الفني (Technical Property) | القيمة المستخرجة (Detected Value) | الملاحظات الهندسية |
| :--- | :--- | :--- |
| **عنوان الآي بي (Target IP)** | `{diagnostic_data.get('target_ip', '127.0.0.1')}` | ثابت في شبكة الـ OT |
| **البروتوكول الصناعي** | **`{diagnostic_data.get('detected_protocol', 'Modbus TCP')}`** | المنفذ المستجيب |
| **المصنع / الفيرموير (Firmware)** | `{diagnostic_data.get('firmware_version', 'Detected / Standard')}` | تم التعرف عليه برمجياً |
| **نوع الموديل (Hardware Type)** | `{diagnostic_data.get('module_type', 'Standard PLC')}` | متوافق مع الحزمة |
| **الرقم التسلسلي (Serial Number)** | `{diagnostic_data.get('serial_number', 'N/A')}` | موثق للأرشفة |
| **مستوى الأمان التشغيلي** | **قراءة فقط صارمة (Strict Read-Only)** | محظور أي أمر كتابة أو إيقاف |
| **متوسط زمن الاستجابة (Latency)** | **`{diagnostic_data.get('avg_latency_ms', 0.0):.2f} ms`** | استجابة القراءة الحية |
| **أقصى سرعة سحب موصى بها** | **`{diagnostic_data.get('recommended_poll_rate_hz', 20)} Hz` ({diagnostic_data.get('recommended_interval_ms', 50)}ms)** | كافية لكشف كل دورات الماكينة |

---

## 2. مصفوفة العناوين المكتشفة (Discovered Memory Map)

{tags_section}

---

## 3. التوصيات الهندسية لأفضل طريقة ربط (AI & OT Recommendations)

{rec_items}

---
*تم إنشاء هذا التقرير تلقائياً بواسطة نظام Machine Network Control.*
"""

        with open(filepath, "w", encoding="utf-8") as f:
            f.write(md_content)

        logger.info(f"Markdown Integration Report generated successfully: {filepath}")
        return filepath

    def generate_tags_config_json(self, diagnostic_data: Dict[str, Any]) -> str:
        """
        Generates a production-ready, machine-readable JSON configuration file
        that can be fed directly to the Polling Engine and SCADA connectors.
        """
        import json
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        station_name = self._sanitize_name(diagnostic_data.get("station_name", "Machine-01"))
        filename = f"tags_config_{station_name}_{timestamp}.json"
        filepath = os.path.join(self.output_dir, filename)

        tags_raw = diagnostic_data.get("tags", [])
        structured_config = {
            "version": "1.0",
            "station_name": station_name,
            "target_ip": diagnostic_data.get("target_ip", "127.0.0.1"),
            "port": diagnostic_data.get("port", 502),
            "protocol": diagnostic_data.get("detected_protocol", "Modbus TCP"),
            "safety_profile": "STRICT_READ_ONLY",
            "intervals_ms": {
                "tier_1_high_speed": 50,
                "tier_2_process": 1000,
                "tier_3_diagnostics": 10000,
            },
            "memory_map": {
                "tags": tags_raw,
                "total_tags": len(tags_raw),
            },
            "poller_ready": True,
            "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        }

        # Write timestamped copy in data/reports
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(structured_config, f, indent=2, ensure_ascii=False)

        # Also write / update current active config in config/active_tags_config.json
        active_config_dir = "config"
        os.makedirs(active_config_dir, exist_ok=True)
        active_filepath = os.path.join(active_config_dir, "active_tags_config.json")
        with open(active_filepath, "w", encoding="utf-8") as f:
            json.dump(structured_config, f, indent=2, ensure_ascii=False)

        logger.info(f"Production Tags JSON Config generated: {filepath} and {active_filepath}")
        return filepath

